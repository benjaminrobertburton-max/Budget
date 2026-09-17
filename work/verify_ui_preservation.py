"""Read-only comparison of a visual redesign against its pre-edit workbook."""
import json
import sys
import zipfile
import xml.etree.ElementTree as ET

import openpyxl

KEYS = {
    "1. Start": ["B5", "B12", "B13", "B14", "B16"],
    "2. Tuesday Review": ["C6", "D15", "D21", "D22", "E24", "E25", "E26", "E27"],
    "3. This Week": ["B3", "B23", "C23", "D23", "E23", "F23"],
    "4. Money Plan": ["B23", "B25", "B26", "B27", "B28"],
    "5. Savings & Debt": ["B5", "B6", "B7", "B9", "D7", "D8", "B17", "B20", "B21", "B39"],
    "6. History": ["F5", "F6", "F7", "G7"],
}


def color(value):
    return str(value) if value is not None else None


def preserved_parts(filename):
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    result = {}
    with zipfile.ZipFile(filename) as archive:
        for name in archive.namelist():
            if name.startswith(("xl/externalLinks/", "xl/pivot", "xl/connection", "xl/vba", "xl/drawings/", "xl/charts/", "xl/comments", "xl/tables/")):
                result[name] = archive.read(name)
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                root = ET.fromstring(archive.read(name))
                for tag in ("dataValidations", "conditionalFormatting", "autoFilter", "sheetProtection", "tableParts"):
                    result[f"{name}:{tag}"] = [ET.tostring(item) for item in root.findall(ns + tag)]
            if name == "xl/workbook.xml":
                root = ET.fromstring(archive.read(name))
                for tag in ("definedNames", "calcPr"):
                    result[f"{name}:{tag}"] = [ET.tostring(item) for item in root.findall(ns + tag)]
    return result


def main():
    before, after = sys.argv[1:3]
    old = openpyxl.load_workbook(before, data_only=False)
    new = openpyxl.load_workbook(after, data_only=False)
    old_results = openpyxl.load_workbook(before, data_only=True)
    new_results = openpyxl.load_workbook(after, data_only=True)
    issues, additions, errors = [], [], []
    formula_count = value_count = input_count = 0
    if old.sheetnames != new.sheetnames:
        issues.append("Sheet order or names changed")
    for name in old.sheetnames:
        a, b = old[name], new[name]
        if a.sheet_state != b.sheet_state:
            issues.append(f"Sheet visibility changed: {name}")
        for row in a:
            for cell in row:
                other = b[cell.coordinate]
                label = f"{name}!{cell.coordinate}"
                if cell.value is not None:
                    if cell.value != other.value or cell.data_type != other.data_type:
                        issues.append(f"Content changed: {label}")
                    value_count += 1
                    formula_count += cell.data_type == "f"
                # Preserve ALL original yellow cells, including blank reserved inputs.
                if cell.fill.fgColor.type == "rgb" and cell.fill.fgColor.rgb[-6:] == "FFF2CC":
                    input_count += 1
                    if (cell.fill.patternType, color(cell.fill.fgColor), color(cell.fill.bgColor), color(cell.font.color)) != (other.fill.patternType, color(other.fill.fgColor), color(other.fill.bgColor), color(other.font.color)):
                        issues.append(f"Input color changed: {label}")
                if old_results[name][cell.coordinate].value != new_results[name][cell.coordinate].value and cell.value is not None:
                    issues.append(f"Calculated/display value changed: {label}")
        for row in b:
            for cell in row:
                if cell.value is not None and a[cell.coordinate].value is None:
                    if cell.hyperlink and cell.hyperlink.location and cell.hyperlink.location.startswith("'"):
                        additions.append(f"{name}!{cell.coordinate}")
                        target, address = cell.hyperlink.location.rsplit("!", 1)
                        target = target[1:-1].replace("''", "'")
                        if target not in new.sheetnames or new[target][address].value is None:
                            issues.append(f"Broken navigation: {name}!{cell.coordinate}")
                    else:
                        issues.append(f"Unexpected new content: {name}!{cell.coordinate}")
        for row in new_results[name]:
            for cell in row:
                if cell.data_type == "e":
                    errors.append(f"{name}!{cell.coordinate}: {cell.value}")
    old_parts, new_parts = preserved_parts(before), preserved_parts(after)
    for key in old_parts.keys() | new_parts.keys():
        if old_parts.get(key) != new_parts.get(key):
            issues.append(f"Native feature changed: {key}")
    output = {
        "sheets": len(old.sheetnames), "preserved_content_cells": value_count,
        "preserved_formulas": formula_count, "preserved_yellow_cells": input_count,
        "new_navigation_links": additions, "formula_errors": errors, "issue_count": len(issues), "issues": issues[:30],
        "key_outputs": {name: {address: new_results[name][address].value for address in addresses} for name, addresses in KEYS.items()},
    }
    print(json.dumps(output, default=str, indent=2))
    return 1 if issues or errors else 0


if __name__ == "__main__":
    sys.exit(main())
