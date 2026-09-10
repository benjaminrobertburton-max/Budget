from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUTPUT = Path(r"C:\Users\Home\Documents\Codex\2026-08-29\help\outputs\01a04fdf-3751-72e2-88f1-daf19b8b9d1d\custom_budget_tracking_solution.docx")
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(89, 99, 112)
BLACK = RGBColor(0, 0, 0)
LIGHT = "F2F4F7"
CALLOUT = "E8EEF5"

def font(run, size=11, color=BLACK, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic

def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)

def margins(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for edge, value in (("top",80),("start",120),("bottom",80),("end",120)):
        node = tc_mar.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_p(doc, text="", style=None, before=0, after=6, size=11, color=BLACK, bold=None, italic=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    font(p.add_run(text), size, color, bold, italic)
    return p

def add_heading(doc, text, level=1):
    return add_p(doc, text, style=f"Heading {level}")

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        font(p.add_run(item), 11)

doc = Document()
section = doc.sections[0]
for side in ("top_margin","bottom_margin","left_margin","right_margin"):
    setattr(section, side, Inches(1))
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
for style_name, size, color, before, after in [
    ("Heading 1",16,BLUE,16,8), ("Heading 2",13,BLUE,12,6), ("Heading 3",12,NAVY,8,4)
]:
    style = doc.styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.1

# Header block: standard_business_brief + memo_masthead.
add_p(doc, "CUSTOM BUDGET SYSTEM - DECISION BRIEF", after=4, size=10.5, color=MUTED, bold=True)
add_p(doc, "A Custom Budget Tracking System That Does Not Depend on One Bank Feed", after=4, size=22, color=NAVY, bold=True)
add_p(doc, "A practical design for weekly budget enforcement across cards, accounts, and pending charges", after=14, size=12, color=MUTED, italic=True)

meta = doc.add_table(rows=3, cols=2)
set_widths(meta, [1.25, 5.25])
for row, (label, detail) in zip(meta.rows, [
    ("Decision", "Build a custom alert-ledger system, not another generic budgeting app or a single aggregator connection."),
    ("Platform", "Private Google Sheet + Google Apps Script + account-activity email alerts + a small mobile exception form."),
    ("Core rule", "Missing data is visible as an exception; it is never treated as zero spending.")
]):
    shade(row.cells[0], LIGHT)
    font(row.cells[0].paragraphs[0].add_run(label), 10.5, NAVY, True)
    font(row.cells[1].paragraphs[0].add_run(detail), 10.5)

add_heading(doc, "The recommendation")
add_p(doc, "Keep the worksheet's targets and logic, but change how actual spending reaches it. Instead of asking an aggregator to provide a complete transaction ledger, collect direct account-activity alert emails from the issuers themselves. Each alert becomes one event in a private ledger. The budget engine then compares pending plus posted activity with the weekly plan.")

callout = doc.add_table(rows=1, cols=1)
set_widths(callout, [6.5])
shade(callout.cell(0,0), CALLOUT)
p = callout.cell(0,0).paragraphs[0]
p.paragraph_format.space_after = Pt(0)
font(p.add_run("Why this is different: "), 11, NAVY, True)
font(p.add_run("a bank aggregator can omit history or pending charges without warning. This system is built to show the gap, preserve the source event, and give you a short exception queue rather than a misleading zero."), 11)

add_heading(doc, "How it works")
steps = [
    ("Direct alerts", "Enable account-activity emails for Wells Fargo, Chase, Citi, Capital One, Discover, and PayPal Credit where available. Use the lowest available transaction threshold and enable purchase, deposit, payment, and transfer alerts."),
    ("Private intake", "A Google Apps Script reads only a dedicated Gmail label, captures the Gmail message ID as a durable event ID, parses each bank's message format, and adds it to the raw ledger. Duplicate emails are ignored."),
    ("Budget rules", "Merchant rules map known transactions to your categories: H-E-B, DoorDash, Starbucks, Apple/Invincibles, fuel, pet, shopping, utilities, subscriptions, and more. Payments and transfers are excluded from spending but tracked as funding actions."),
    ("Weekly scorecard", "A single Tuesday-Monday screen shows Target, Posted, Pending, Committed, Remaining, and Needs Review. Committed spending is the safe-to-spend input."),
    ("Exception form", "A small phone form handles cash, a push-only alert, a parser failure, a tip correction, or a transfer confirmation. It takes seconds and is the only regular manual entry."),
    ("Tuesday reconciliation", "Review only the exception queue and confirm current card/checking balances. Stale or unmatched pending events are flagged before a week is closed.")
]
for label, detail in steps:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.1
    font(p.add_run(label + ": "), 11, NAVY, True)
    font(p.add_run(detail), 11)

add_heading(doc, "What the system contains")
tbl = doc.add_table(rows=1, cols=3)
set_widths(tbl, [1.55, 2.25, 2.7])
for cell, text in zip(tbl.rows[0].cells, ["Component", "What it does", "Why it matters"]):
    shade(cell, LIGHT)
    font(cell.paragraphs[0].add_run(text), 10.5, NAVY, True)
for vals in [
    ("Raw alert ledger", "Stores event ID, received time, institution, account, merchant, amount, source, and status.", "Preserves every source event and prevents silent duplicates."),
    ("Rules & exceptions", "Stores merchant mappings, transfer pairing, and unknown-event review.", "Improves with your actual habits while keeping ambiguity visible."),
    ("Weekly scorecard", "Calculates Target, Posted, Pending, Committed, and Remaining.", "Answers whether you are still within the budget before next Tuesday."),
    ("Funding check", "Tracks rent, Fidelity, cruise, tuition, and Wealthfront actions separately.", "Prevents transfers and card payments from being double-counted as spending."),
    ("Weekly history", "Keeps one row per Tuesday-Monday week and chart-ready totals.", "Shows trends without a monthly statement rebuild or 52 worksheet tabs.")
]:
    cells = tbl.add_row().cells
    for cell, text in zip(cells, vals):
        font(cell.paragraphs[0].add_run(text), 10.2)

add_heading(doc, "Your actual weekly workflow")
add_bullets(doc, [
    "One-time: enable alerts and route their emails to a dedicated Gmail label. Test every account with a real transaction or a known recurring charge before trusting its parser.",
    "Most days: do nothing. The issuer alert creates the ledger event. If an alert never arrives or a charge is cash/push-only, add it through the phone exception form.",
    "Tuesday: open the Review queue, resolve only exceptions, confirm account balances, and let the scorecard close the week. This is a review, not a statement-rebuild project."
])

add_heading(doc, "Limits and safeguards")
add_bullets(doc, [
    "No system can record a charge that the bank neither sends as an email nor exposes elsewhere. This design flags coverage gaps instead of silently understating spending.",
    "Email contents are financial data. Use a dedicated Gmail label, keep two-factor authentication enabled, and authorize the script only for your own Gmail and spreadsheet.",
    "Pending authorizations can settle for a different amount or reverse. The scorecard treats them as committed for safe-to-spend; reconciliation resolves the final difference.",
    "A transfer can create two alerts. Pairing rules prevent it from appearing as both spending and income."
])

add_heading(doc, "Why this is the right reset")
compare = doc.add_table(rows=1, cols=3)
set_widths(compare, [2.05, 2.2, 2.25])
for cell, text in zip(compare.rows[0].cells, ["Approach", "What failed", "Custom-system response"]):
    shade(cell, LIGHT)
    font(cell.paragraphs[0].add_run(text), 10.5, NAVY, True)
for vals in [
    ("Aggregator / Empower", "Partial history and incomplete pending activity.", "Uses issuer-originated alerts and a visible exception queue."),
    ("Manual weekly totals", "Too many charges across too many accounts.", "Automates normal events; only misses and corrections need manual entry."),
    ("Generic budgeting app", "May inherit the same feed limitations and cannot reflect your custom logic.", "Keeps your categories, savings rules, tuition plan, and transfer treatment.")
]:
    cells = compare.add_row().cells
    for cell, text in zip(cells, vals):
        font(cell.paragraphs[0].add_run(text), 10.2)

add_heading(doc, "What I need before building it")
add_p(doc, "The next step is not another data test. It is approval to build this private Google Sheets system and confirmation that you can use Gmail for a dedicated financial-alert label. Once approved, we start with Wells Fargo and Chase, test those parsers for two weeks, then add the other active spending accounts.")

add_heading(doc, "Evidence consulted")
for name, url in [
    ("Chase Account Alerts", "https://www.chase.com/personal/mobile-online-banking/login-alerts"),
    ("Wells Fargo Alerts Questions", "https://www.wellsfargo.com/help/online-banking/alerts-faqs/"),
    ("Capital One purchase notifications", "https://www.capitalone.com/learn-grow/privacy-security/protect-digital-identity/"),
    ("Google Gmail filters", "https://support.google.com/mail/answer/6579"),
    ("Google Apps Script GmailApp", "https://developers.google.com/apps-script/reference/gmail/gmail-app"),
    ("Google Sheets appendRow", "https://developers.google.com/apps-script/reference/spreadsheet/sheet"),
    ("Empower pending-transaction limitation", "https://support-personalwealth.empower.com/hc/en-us/articles/201170060-Why-can-t-I-see-my-pending-transactions")
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    font(p.add_run(name + ": "), 9.5, BLACK, True)
    font(p.add_run(url), 9.5, BLUE)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(footer.add_run("Custom budget tracking solution | 2026-08-30"), 9, MUTED)

doc.save(OUTPUT)
print(str(OUTPUT))

